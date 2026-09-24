# Отчёт по аудиту: уровень существенности и признаки «положительного» баланса АО «Транспневматика»
import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

OUT = sys.argv[1] if len(sys.argv) > 1 else str(
    Path(__file__).resolve().parent.parent / "Сомов Н.В._3524Б6ЭКэп1_Аудит_задание на 3 мая.docx")
NB = " "


def fmt(x, dec=0, sign=False):
    s = f"{abs(x):,.{dec}f}".replace(",", NB).replace(".", ",")
    if x < 0:
        s = "−" + s
    elif sign and x > 0:
        s = "+" + s
    return s


def nb(t):
    # неразрывные пробелы, чтобы «тыс. руб.», «%» и «№» не отрывались от чисел
    return (t.replace(" тыс. руб", NB + "тыс." + NB + "руб").replace(" %", NB + "%")
             .replace("№ ", "№" + NB).replace(", 4, 5", "," + NB + "4," + NB + "5"))


def pct(x, dec=2, sign=False):
    return fmt(x, dec, sign)


# ---------- исходные данные, тыс. руб. ----------
profit_bt = 197_091          # 2300
revenue = 3_139_206          # 2110
balance = 3_500_752          # 1600
equity = 2_731_361           # 1300
cost = 2_662_114             # 2120
selling = 137_770            # 2210
admin = 0                    # 2220, прочерк
costs = cost + selling + admin

inv = {"Сырьё и материалы": 454_350, "Незавершённое производство": 151_046,
       "Готовая продукция": 370_697, "Товары отгруженные": 0,
       "Расходы будущих периодов": 3_893}
inventory = 979_986
assert sum(inv.values()) == inventory

base = [
    ("Прибыль до налогообложения", "стр. 2300 ОФР", profit_bt, 5),
    ("Выручка (без НДС)", "стр. 2110 ОФР", revenue, 2),
    ("Валюта баланса", "стр. 1600 ББ", balance, 2),
    ("Собственный капитал", "стр. 1300 ББ", equity, 10),
    ("Общие затраты", "стр. 2120 + 2210 + 2220 ОФР", costs, 2),
]
vals = [v * p / 100 for _, _, v, p in base]
i_max, i_min = vals.index(max(vals)), vals.index(min(vals))
kept = [v for i, v in enumerate(vals) if i not in (i_max, i_min)]
avg5 = sum(vals) / 5
avg3 = sum(kept) / 3
SBFO = 60_000
dev_round = (avg3 - SBFO) / avg3 * 100
assert dev_round <= 20

inv_share = inventory / balance
m_inv = SBFO * inv_share
fg = inv["Готовая продукция"]
fg_share_inv = fg / inventory
m_fg = m_inv * fg_share_inv

# ---------- задание 2 ----------
b13 = dict(nca=1_533_750, ca=1_838_703, assets=3_372_453, ar=490_512, cash=507_757,
           eq=2_722_277, np=200_305, lt=431_036, st=219_140, ap=219_140)
b14 = dict(nca=1_643_615, ca=1_857_137, assets=3_500_752, ar=756_782, cash=52_635,
           eq=2_731_361, np=127_358, lt=487_210, st=282_181, ap=232_181)
for b in (b13, b14):
    b["debt"] = b["lt"] + b["st"]
    b["sos"] = (b["eq"] - b["nca"]) / b["ca"] * 100

rows2 = [
    ("Внеоборотные активы (стр. 1100)", "nca"),
    ("Оборотные активы (стр. 1200)", "ca"),
    ("Активы (имущество) всего (стр. 1600)", "assets"),
    ("Дебиторская задолженность (стр. 1230)", "ar"),
    ("Денежные средства (стр. 1250)", "cash"),
    ("Собственный капитал (стр. 1300)", "eq"),
    ("Прибыль (чистая прибыль, стр. 2400)", "np"),
    ("Заёмный капитал (стр. 1400 + 1500)", "debt"),
    ("Кредиторская задолженность (стр. 1520)", "ap"),
    ("Доля собственных средств в оборотных активах, %", "sos"),
]


def g(k):
    a, b = b13[k], b14[k]
    return a, b, b - a, b / a * 100, b / a * 100 - 100


# ---------- оформление ----------
doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21)
sec.left_margin, sec.right_margin = Cm(3), Cm(1.5)
sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)

st = doc.styles["Normal"]
st.font.name = "Times New Roman"
st.font.size = Pt(14)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
pf = st.paragraph_format
pf.space_before = pf.space_after = Pt(0)
pf.line_spacing = 1.5


def para(text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, size=None, indent=True,
         before=0, after=0, italic=False, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    f = p.paragraph_format
    f.first_line_indent = Cm(1.25) if indent else Cm(0)
    f.space_before, f.space_after = Pt(before), Pt(after)
    if keep:
        f.keep_with_next = True
    if text:
        r = p.add_run(nb(text))
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
    return p


def rich(parts, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=True):
    # parts: [(текст, bold)]
    p = para(align=align, indent=indent)
    for t, b in parts:
        p.add_run(nb(t)).bold = b
    return p


def formula(text):
    return para(text, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)


def heading(text, before=12, after=6):
    return para(text, align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, indent=False,
                before=before, after=after, keep=True)


def caption(text):
    p = para(text, align=WD_ALIGN_PARAGRAPH.LEFT, indent=False, before=6, keep=True)
    return p


def shade(cell, color="E7E6E6"):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), color)
    tcPr.append(sh)


def table(header, rows, widths, align_cols=None, bold_rows=(), size=12, shade_rows=(), center_rows=()):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    t._tbl.tblPr.append(lay)
    align_cols = align_cols or {}
    for ri, row in enumerate([header] + rows):
        for ci, val in enumerate(row):
            c = t.cell(ri, ci)
            c.width = Cm(widths[ci])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.first_line_indent = Cm(0)
            if ri == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if (ri - 1) in center_rows
                               else align_cols.get(ci, WD_ALIGN_PARAGRAPH.RIGHT))
            r = p.add_run(str(val) if ri == 0 else nb(str(val)))
            if ri < len(rows):
                p.paragraph_format.keep_with_next = True
            r.font.size = Pt(size)
            r.bold = ri == 0 or (ri - 1) in bold_rows
            if ri == 0:
                shade(c)
            elif (ri - 1) in shade_rows:
                shade(c, "F2F2F2")
    trPr = t.rows[0]._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trPr.append(h)
    for row in t.rows:
        tr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        cs.set(qn("w:val"), "true")
        tr.append(cs)
    doc.add_paragraph().paragraph_format.line_spacing = 1.0
    return t


L, C = WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER

# ---------- титул ----------
for line in ["МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ", "РОССИЙСКОЙ ФЕДЕРАЦИИ",
             "Федеральное государственное автономное образовательное учреждение",
             "высшего образования",
             "«Национальный исследовательский Нижегородский государственный",
             "университет им. Н.И. Лобачевского»"]:
    para(line, align=C, indent=False, size=12)
para("Институт экономики и предпринимательства", align=C, indent=False, size=12, before=6)
para("Кафедра финансов и кредита", align=C, indent=False, size=12)
para("ПРАКТИЧЕСКОЕ ЗАДАНИЕ", align=C, indent=False, bold=True, before=16)
para("по дисциплине «Аудит»", align=C, indent=False)
para("Уровень существенности и признаки «положительного» баланса", align=C, indent=False)
para("АО «Транспневматика» (отчётность за 2014 год)", align=C, indent=False)
for i, line in enumerate(["Выполнил: студент 3 курса, группа 3524Б6ЭКэп1,",
                          "направление 38.03.01 «Экономика»,",
                          "профиль «Экономика, международный бизнес",
                          "и предпринимательство»",
                          "Сомов Никита Владимирович"]):
    para(line, align=WD_ALIGN_PARAGRAPH.RIGHT, indent=False, size=13, before=12 if i == 0 else 0)
para("Нижний Новгород, 2026", align=C, indent=False, before=12, after=10)

# ---------- задание 1 ----------
heading("Задание 1. Уровень существенности и удельная существенность", before=6)
para("Исходные данные — бухгалтерский баланс АО «Транспневматика» на 31.12.2014, "
     "отчёт о финансовых результатах за 2014 год и пояснения по строке «Запасы». "
     "Все суммы — в тыс. руб.")

heading("1.1. СБФО по базовым отчётным показателям")
para("Значение для нахождения уровня существенности по каждому базовому показателю "
     "равно произведению показателя на установленную долю: 5 % для прибыли "
     "до налогообложения, 2 % для выручки, валюты баланса и общих затрат, "
     "10 % для собственного капитала.")
para("Управленческие расходы за 2014 год по стр. 2220 не показаны (прочерк), поэтому "
     f"общие затраты равны сумме себестоимости продаж и коммерческих расходов: "
     f"{fmt(cost)} + {fmt(selling)} = {fmt(costs)} тыс. руб.")

caption("Таблица 1 — Расчёт значений для нахождения уровня существенности")
rows = []
for i, ((name, src, v, p), val) in enumerate(zip(base, vals)):
    mark = ""
    if i == i_max:
        mark = " (наиб., искл.)"
    elif i == i_min:
        mark = " (наим., искл.)"
    rows.append([str(i + 1), name + mark, src, fmt(v), f"{p}", fmt(val, 1)])
table(["№", "Базовый показатель", "Источник", "Значение, тыс. руб.", "Доля, %",
       "Значение для УС, тыс. руб."],
      rows, [0.8, 4.6, 3.6, 2.8, 1.5, 3.2],
      align_cols={0: C, 1: L, 2: L, 4: C}, shade_rows=(i_max, i_min))

para(f"Наибольшее значение даёт собственный капитал ({fmt(vals[i_max], 1)} тыс. руб.), "
     f"наименьшее — прибыль до налогообложения ({fmt(vals[i_min], 1)} тыс. руб.). "
     f"Они отличаются от среднего по пяти показателям ({fmt(avg5, 1)} тыс. руб.) "
     f"на {pct((vals[i_max] - avg5) / avg5 * 100, 1, True)} % и "
     f"{pct((vals[i_min] - avg5) / avg5 * 100, 1)} % и по условию исключаются. "
     "Среднее по оставшимся трём показателям:")
formula(f"({fmt(kept[0], 1)} + {fmt(kept[1], 1)} + {fmt(kept[2], 1)}) / 3 = "
        f"{fmt(avg3, 1)} тыс. руб.")
para(f"Допустимое округление — не более 20 %, то есть в пределах от "
     f"{fmt(avg3 * 0.8, 1)} до {fmt(avg3 * 1.2, 1)} тыс. руб. Среднее округляем "
     f"в меньшую сторону до {fmt(SBFO)} тыс. руб.: более низкий порог делает "
     "проверку строже. Отклонение от неокруглённого значения:")
formula(f"({fmt(avg3, 1)} − {fmt(SBFO)}) / {fmt(avg3, 1)} × 100 % = "
        f"{pct(dev_round)} % < 20 %.")
rich([("Вывод: ", True),
      (f"уровень существенности (СБФО) АО «Транспневматика» за 2014 год — "
       f"{fmt(SBFO)} тыс. руб., это {pct(SBFO / balance * 100)} % валюты баланса. "
       "Искажения отчётности в меньшей сумме не влияют на решения пользователей.", False)])

heading("1.2. Удельная существенность по статье «Запасы»")
para("Общий уровень существенности распределяется между статьями баланса "
     "пропорционально их удельному весу в валюте баланса:")
formula("УС(статьи) = СБФО × Статья / Валюта баланса.")
para(f"Удельный вес запасов (стр. 1210) в валюте баланса (стр. 1600):")
formula(f"{fmt(inventory)} / {fmt(balance)} × 100 % = {pct(inv_share * 100)} %.")
formula(f"УС(запасы) = {fmt(SBFO)} × {fmt(inv_share, 4)} = {fmt(m_inv, 1)} тыс. руб.")
rich([("Вывод: ", True),
      (f"удельная существенность по статье «Запасы» — {fmt(m_inv, 1)} тыс. руб. "
       f"(≈ {fmt(round(m_inv))} тыс. руб.). Выявленные по запасам ошибки меньше этой "
       "суммы аудитор может признать несущественными.", False)])

heading("1.3. Удельная существенность для объекта аудита «Готовая продукция»")
para("В балансе готовая продукция входит в строку «Запасы», её остаток берём "
     f"из пояснений: стр. 5403, гр. 8 — {fmt(fg)} тыс. руб. на 31.12.2014. "
     "Удельная существенность статьи «Запасы» распределяется между группами "
     "запасов пропорционально их доле:")
formula(f"Доля готовой продукции в запасах = {fmt(fg)} / {fmt(inventory)} × 100 % = "
        f"{pct(fg_share_inv * 100)} %.")
formula(f"УС(готовая продукция) = {fmt(m_inv, 1)} × {fmt(fg_share_inv, 4)} = "
        f"{fmt(m_fg, 1)} тыс. руб.")
para("Проверка через валюту баланса даёт то же значение:")
formula(f"{fmt(SBFO)} × {fmt(fg)} / {fmt(balance)} = {fmt(SBFO * fg / balance, 1)} тыс. руб.")

caption("Таблица 2 — Удельная существенность по группам запасов на 31.12.2014")
rows = []
for name, v in inv.items():
    rows.append([name, fmt(v), pct(v / inventory * 100), pct(v / balance * 100),
                 fmt(m_inv * v / inventory, 1)])
rows.append(["Запасы — всего", fmt(inventory), "100,00", pct(inv_share * 100), fmt(m_inv, 1)])
fg_idx = list(inv).index("Готовая продукция")
table(["Группа запасов", "Сумма, тыс. руб.", "Доля в запасах, %",
       "Доля в валюте баланса, %", "Удельная существенность, тыс. руб."],
      rows, [4.6, 2.6, 2.4, 2.8, 4.1], align_cols={0: L},
      bold_rows=(len(rows) - 1, fg_idx), shade_rows=(fg_idx,))

rich([("Вывод: ", True),
      (f"удельная существенность для объекта аудита «Готовая продукция» — "
       f"{fmt(m_fg, 1)} тыс. руб. (≈ {fmt(round(m_fg))} тыс. руб.). ", False),
      ("Объект требует внимания: за 2014 год остаток готовой продукции вырос "
       f"с {fmt(214_046)} до {fmt(fg)} тыс. руб. (на "
       f"{pct((fg / 214_046 - 1) * 100, 1)} %), а выручка снизилась на "
       f"{pct((1 - 3_139_206 / 3_376_778) * 100, 1)} %. Это повышает риск "
       "искажения оценки готовой продукции (затоваривание, неликвиды, "
       "списание до чистой стоимости продажи).", False)])

# ---------- задание 2 ----------
heading("Задание 2. Оценка признаков «положительного» баланса", before=12)
para("Для оценки заполнена аналитическая таблица по данным баланса на 31.12.2013 "
     "и 31.12.2014 и отчёта о финансовых результатах за 2013 и 2014 годы. "
     "Заёмный капитал — сумма долгосрочных и краткосрочных обязательств. Доля "
     "собственных средств в оборотных активах рассчитана как (стр. 1300 − стр. 1100) / "
     "стр. 1200 × 100 %; для неё абсолютное отклонение дано в процентных пунктах.")

caption("Таблица 3 — Динамика показателей АО «Транспневматика»")
rows = []
for i, (name, k) in enumerate(rows2):
    a, b, d, tr, tp = g(k)
    is_pct = k == "sos"
    dec = 2 if is_pct else 0
    rows.append([f"{i + 1}. {name}", fmt(a, dec), fmt(b, dec),
                 fmt(d, dec, True) + (" п.п." if is_pct else ""),
                 pct(tr), pct(tp, 2, True)])
table(["Показатели", "на 31 дек. 2013 г.", "на 31 дек. 2014 г.",
       "Абсолютное отклонение (гр. 3 − гр. 2), тыс. руб.",
       "Темп роста (гр. 3 / гр. 2 × 100), %", "Темп прироста (гр. 5 − 100), %"],
      [["1", "2", "3", "4", "5", "6"]] + rows,
      [4.6, 2.4, 2.4, 2.6, 2.3, 2.2], align_cols={0: L}, size=11, center_rows=(0,))

sos13, sos14 = b13["sos"], b14["sos"]
tp = {k: g(k)[4] for _, k in rows2}
tr = {k: g(k)[3] for _, k in rows2}
checks = [
    ("Валюта баланса на конец периода больше, чем на начало",
     f"{fmt(b14['assets'])} > {fmt(b13['assets'])}; прирост {pct(tp['assets'], 2, True)} %", True),
    ("Темп прироста оборотных активов выше темпа прироста внеоборотных",
     f"{pct(tp['ca'], 2, True)} % < {pct(tp['nca'], 2, True)} %", False),
    ("Собственный капитал превышает заёмный",
     f"{fmt(b14['eq'])} > {fmt(b14['debt'])} (в {fmt(b14['eq'] / b14['debt'], 2)} раза)", True),
    ("Темп роста собственного капитала выше темпа роста заёмного",
     f"{pct(tr['eq'])} % < {pct(tr['debt'])} %", False),
    ("Темпы прироста дебиторской и кредиторской задолженности примерно одинаковы",
     f"ДЗ {pct(tp['ar'], 2, True)} %, КЗ {pct(tp['ap'], 2, True)} %", False),
    ("Доля собственных средств в оборотных активах более 10 %",
     f"{pct(sos14)} % (2013 г. — {pct(sos13)} %)", True),
    ("В балансе нет непокрытого убытка",
     f"нераспределённая прибыль {fmt(2_635_111)}", True),
]
caption("Таблица 4 — Соблюдение признаков «положительного» баланса на 31.12.2014")
table(["№", "Признак «положительного» баланса", "Фактические данные", "Выполнение"],
      [[str(i + 1), n, f, "да" if ok else "нет"] for i, (n, f, ok) in enumerate(checks)],
      [0.8, 6.8, 5.9, 3.0], align_cols={0: C, 1: L, 2: L, 3: C}, size=12)

ok_n = sum(1 for *_, ok in checks if ok)
para(f"Из семи признаков выполнены {ok_n} (№ 1, 3, 6, 7), не выполнены "
     f"{len(checks) - ok_n} (№ 2, 4, 5).")
para("Выполненные признаки говорят о высокой финансовой устойчивости: собственный "
     f"капитал составляет {pct(b14['eq'] / b14['assets'] * 100, 1)} % пассива, "
     f"собственными средствами покрыто {pct(sos14, 1)} % оборотных активов при нормативе "
     "10 %, убытков нет, валюта баланса растёт.")
para("Невыполненные признаки и другие показатели таблицы 3 показывают ухудшение "
     "за 2014 год:")
for t in [
    f"прирост вложений во внеоборотные активы ({pct(tp['nca'], 2, True)} %) "
    f"в семь раз обгоняет прирост оборотных ({pct(tp['ca'], 2, True)} %);",
    f"заёмный капитал растёт быстрее собственного ({pct(tp['debt'], 2, True)} % "
    f"против {pct(tp['eq'], 2, True)} %), появились краткосрочные займы "
    f"{fmt(50_000)} тыс. руб.;",
    f"дебиторская задолженность выросла на {pct(tp['ar'], 1)} %, кредиторская — "
    f"только на {pct(tp['ap'], 1)} %: средства отвлекаются в расчёты с покупателями;",
    f"денежные средства сократились с {fmt(b13['cash'])} до {fmt(b14['cash'])} тыс. руб. "
    f"(в {fmt(b13['cash'] / b14['cash'], 1)} раза), чистая прибыль — "
    f"на {pct(-tp['np'], 1)} %, выручка — на {pct((1 - 3_139_206 / 3_376_778) * 100, 1)} %.",
]:
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(nb(t))
rich([("Вывод: ", True),
      ("отчётность АО «Транспневматика» на 31.12.2014 соблюдает признаки "
       "«положительного» баланса лишь частично. Структура капитала устойчивая, "
       "но динамика отрицательная. Для аудитора это зоны повышенного риска "
       "существенного искажения: дебиторская задолженность (реальность долгов, "
       "резерв по сомнительным долгам), запасы и готовая продукция (оценка, "
       "неликвиды), денежные средства и прочие расходы, которые выросли "
       f"с {fmt(363_474)} до {fmt(687_959)} тыс. руб. Отдельно следует проверить "
       "сопоставимость управленческих расходов: в 2013 году они показаны "
       f"в сумме {fmt(432_967)} тыс. руб., в 2014 году — прочерк.", False)])

doc.save(OUT)
print("saved", OUT)
print(f"avg5={avg5:.2f} avg3={avg3:.2f} dev={dev_round:.2f} m_inv={m_inv:.2f} m_fg={m_fg:.2f}")
for name, k in rows2:
    print(name, [round(x, 2) for x in g(k)])
